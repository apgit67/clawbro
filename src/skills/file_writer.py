"""
skills/file_writer.py
---------------------
Generates file content with Claude and writes it to disk in the
requested format: .txt, .md, .csv, .json, .html, .docx (Word), or .pdf.

Optional dependencies (install to unlock the format):
  Word (.docx) : pip install python-docx
  PDF  (.pdf)  : pip install fpdf2
"""

from __future__ import annotations

import json
import os
import re
from datetime import datetime
from pathlib import Path

from core.context import ConversationContext, SkillResult
from skills.base import SkillBase

# ---------------------------------------------------------------------------
# Format detection
# ---------------------------------------------------------------------------

_FORMAT_PATTERNS: dict[str, list[str]] = {
    "docx": [r"\bdocx\b", r"word\s*doc", r"\.docx"],
    "pdf":  [r"\bpdf\b", r"\.pdf"],
    "md":   [r"\bmarkdown\b", r"\bmark\s*down\b", r"\.md\b"],
    "csv":  [r"\bcsv\b", r"\.csv"],
    "json": [r"\bjson\b", r"\.json"],
    "html": [r"\bhtml\b", r"\.html"],
    "txt":  [r"\b(txt|text\s*file|notepad|plain\s*text)\b", r"\.txt"],
}

# More specific formats are checked before the plain-text fallback
_FORMAT_ORDER = ["docx", "pdf", "md", "csv", "json", "html", "txt"]

# "save to ~/docs/report.pdf", "save as report.pdf", "output to/as notes.txt"
_EXPLICIT_PATH_RE = re.compile(
    r"""
    (?:save\s+(?:it\s+)?(?:to|as|in)\s+
    |  output\s+(?:to|as)\s+
    |  write\s+(?:it\s+)?(?:to|as)\s+
    )
    ([^\s"'<>|?*]+\.[a-zA-Z]{2,5})
    """,
    re.IGNORECASE | re.VERBOSE,
)

# "call it report", "filename: notes", "name the file summary"
_FILENAME_RE = re.compile(
    r"""
    (?:call\s+it\s+
    |  name\s+(?:it\s+|the\s+file\s+)?
    |  filename[:\s]+
    |  file\s+name[:\s]+
    )
    ["']?([^\s"'<>|?*\.]+)["']?
    """,
    re.IGNORECASE | re.VERBOSE,
)

_DEFAULT_OUTPUT_DIR = Path("~/.clawbro/files").expanduser()


class UnsafePathError(ValueError):
    """Raised when a requested output path escapes the allowed output directory."""


def _contain_path(candidate: Path) -> Path:
    """Resolve a candidate path and ensure it stays inside _DEFAULT_OUTPUT_DIR.

    A user message can request an explicit path; without this check that path
    could be absolute or contain '..' segments, giving an arbitrary-file-write
    primitive. We collapse the candidate against the output dir and reject
    anything that resolves outside it.
    """
    base = _DEFAULT_OUTPUT_DIR.resolve()
    # Absolute paths and traversal are reinterpreted relative to the base:
    # only the final filename is honored, the rest of any path is dropped.
    resolved = (base / candidate.name).resolve()
    if resolved != base and not resolved.is_relative_to(base):
        raise UnsafePathError(
            f"Refusing to write outside {base}: {candidate}"
        )
    return resolved


def _detect_format(message: str) -> str:
    lowered = message.lower()
    for fmt in _FORMAT_ORDER:
        for pattern in _FORMAT_PATTERNS[fmt]:
            if re.search(pattern, lowered):
                return fmt
    return "txt"


def _detect_output_path(message: str, fmt: str) -> Path:
    # Explicit path: honor only its filename, contained within the output dir.
    m = _EXPLICIT_PATH_RE.search(message)
    if m:
        return _contain_path(Path(m.group(1)))

    # Named file without full path
    m = _FILENAME_RE.search(message)
    if m:
        stem = m.group(1)
        return _contain_path(_DEFAULT_OUTPUT_DIR / f"{stem}.{fmt}")

    # Auto-generated timestamped filename
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _contain_path(_DEFAULT_OUTPUT_DIR / f"clawbro_{ts}.{fmt}")


# ---------------------------------------------------------------------------
# Format writers
# ---------------------------------------------------------------------------

def _write_txt(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def _write_md(path: Path, content: str) -> None:
    _write_txt(path, content)


def _write_csv(path: Path, content: str) -> None:
    _write_txt(path, content)


def _write_json(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        data = json.loads(content)
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    except (json.JSONDecodeError, ValueError):
        # Content isn't valid JSON — write as-is
        _write_txt(path, content)


def _write_html(path: Path, content: str) -> None:
    if not content.strip().lower().startswith("<!doctype") and "<html" not in content.lower():
        content = (
            "<!DOCTYPE html>\n"
            '<html lang="en">\n'
            '<head><meta charset="UTF-8"><title>ClawBro Output</title></head>\n'
            "<body>\n<pre>\n"
            + content
            + "\n</pre>\n</body>\n</html>"
        )
    _write_txt(path, content)


def _write_docx(path: Path, content: str) -> None:
    try:
        from docx import Document  # type: ignore[import]
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in content.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


def _write_pdf(path: Path, content: str) -> None:
    try:
        from fpdf import FPDF  # type: ignore[import]
    except ImportError as exc:
        raise RuntimeError(
            "fpdf2 is not installed. Run: pip install fpdf2"
        ) from exc
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in content.splitlines():
        # fpdf2 requires latin-1 safe text; replace unsupported chars gracefully
        safe_line = line.encode("latin-1", errors="replace").decode("latin-1")
        pdf.cell(0, 7, safe_line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


_WRITERS = {
    "txt":  _write_txt,
    "md":   _write_md,
    "csv":  _write_csv,
    "json": _write_json,
    "html": _write_html,
    "docx": _write_docx,
    "pdf":  _write_pdf,
}

# ---------------------------------------------------------------------------
# Claude system prompt
# ---------------------------------------------------------------------------

_SYSTEM_PROMPT = """\
You are a precise document writer. Generate the content for the file the user requests.

Rules:
- Output ONLY the file content — no preamble, no "Here is the file:", no explanation.
- Do NOT wrap output in markdown code fences unless the target format is Markdown.
- For CSV: include a header row; use comma delimiters.
- For JSON: output valid, well-formatted JSON only.
- For HTML: output a complete HTML document with proper structure.
- For plain text, Markdown, Word, or PDF: write clear, well-structured prose or
  structured content appropriate to the request.
"""


# ---------------------------------------------------------------------------
# Skill class
# ---------------------------------------------------------------------------

class FileWriterSkill(SkillBase):
    """Generates content with Claude and saves it to a file in the requested format."""

    name = "file_writer"
    description = (
        "Generates content and writes it to a file. Supports .txt, .md, .pdf, "
        ".docx (Word), .csv, .json, and .html. Saves to ~/.clawbro/files/ by default."
    )
    version = "1.0.0"
    trigger_patterns = [
        r"write.*file",
        r"create.*file",
        r"save.*(?:as|to\s+a?\s*file)",
        r"export.*(?:as|to)",
        r"generate.*(?:file|document|doc\b)",
        r"make.*(?:file|document|doc\b)",
        r"\bword\s*doc",
        r"\bdocx\b",
        r"(?:write|create|save|export|generate).*pdf",
        r"text\s*file",
        r"save.*notepad",
        r"\.txt\b",
        r"\.md\b",
        r"\.csv\b",
        r"\.json\b",
        r"\.html\b",
    ]

    def handle(self, context: ConversationContext) -> SkillResult:
        try:
            user_text = context.message.text
            fmt = _detect_format(user_text)
            out_path = _detect_output_path(user_text, fmt)

            messages = [
                {"role": "user", "content": _SYSTEM_PROMPT},
                *context.history,
                {"role": "user", "content": user_text},
            ]
            content = context.claude.complete(messages)

            if not content or not content.strip():
                return self._error_result("Claude returned empty content — nothing to write.")

            _WRITERS[fmt](out_path, content)

            size = out_path.stat().st_size
            preview = content[:500] + ("..." if len(content) > 500 else "")

            return SkillResult(
                text=(
                    f"File written successfully.\n"
                    f"  Path:   {out_path}\n"
                    f"  Format: {fmt.upper()}\n"
                    f"  Size:   {size:,} bytes\n\n"
                    f"--- Preview ---\n{preview}"
                ),
                skill_name=self.name,
                success=True,
                metadata={"format": fmt, "path": str(out_path), "size": size},
            )
        except Exception as exc:  # noqa: BLE001
            return self._error_result(str(exc))
